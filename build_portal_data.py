import openpyxl
import json

wb = openpyxl.load_workbook('Gabarito_Maratona_30_Questoes.xlsx', data_only=True)
ws = wb['Todas as Turmas']

turma_blocks = [
    ("6º Ano", 4, 18),
    ("7º Ano", 20, 42),
    ("8º Ano", 44, 53),
    ("9º Ano", 55, 66),
    ("1ª Série EM", 68, 78),
    ("2ª Série EM", 80, 92),
    ("3ª Série EM", 94, 108),
]

all_students = []
turmas_data = {}

for turma_name, start_r, end_r in turma_blocks:
    gab_row = start_r + 2
    gabarito = [str(ws.cell(gab_row, c).value or '').strip().upper() for c in range(3, 33)]
    
    turma_students = []
    for r in range(start_r + 3, end_r + 1):
        num = ws.cell(r, 1).value
        name = ws.cell(r, 2).value
        if not name:
            continue
        answers = [str(ws.cell(r, c).value or '').strip().upper() for c in range(3, 33)]
        present = any(a != '' for a in answers)
        
        r_score = 0
        l_score = 0
        details = []
        for i in range(30):
            q_num = i + 1
            correct_ans = gabarito[i]
            stu_ans = answers[i]
            is_annulled = (turma_name in ["8º Ano", "9º Ano"] and q_num == 4)
            is_correct = False
            
            if present:
                if is_annulled:
                    is_correct = True
                elif stu_ans and correct_ans and stu_ans == correct_ans:
                    is_correct = True
                    
            if is_correct:
                if q_num <= 20:
                    r_score += 1
                else:
                    l_score += 1
                    
            details.append({
                "q": q_num,
                "domain": "Reading" if q_num <= 20 else "Listening",
                "student_ans": stu_ans if stu_ans else "-",
                "correct_ans": "*" if is_annulled else correct_ans,
                "is_correct": is_correct,
                "is_annulled": is_annulled
            })
            
        total_score = r_score + l_score
        pct_total = round((total_score / 30.0) * 100, 1) if present else 0.0
        pct_r = round((r_score / 20.0) * 100, 1) if present else 0.0
        pct_l = round((l_score / 10.0) * 100, 1) if present else 0.0
        
        stu_obj = {
            "id": f"{turma_name.replace(' ', '_').lower()}_{num}",
            "num": num,
            "name": str(name).strip(),
            "turma": turma_name,
            "present": present,
            "total_score": total_score,
            "pct_total": pct_total,
            "reading_score": r_score,
            "pct_reading": pct_r,
            "listening_score": l_score,
            "pct_listening": pct_l,
            "details": details
        }
        turma_students.append(stu_obj)
        all_students.append(stu_obj)
        
    # Rank within turma
    present_turma = [s for s in turma_students if s["present"]]
    present_turma.sort(key=lambda s: (s["total_score"], s["reading_score"], s["listening_score"]), reverse=True)
    
    current_rank = 1
    for i, s in enumerate(present_turma):
        if i > 0:
            prev = present_turma[i-1]
            if (s["total_score"], s["reading_score"], s["listening_score"]) < (prev["total_score"], prev["reading_score"], prev["listening_score"]):
                current_rank = i + 1
        s["rank_turma"] = current_rank
        
    for s in turma_students:
        if not s["present"]:
            s["rank_turma"] = "-"
            
    turmas_data[turma_name] = {
        "gabarito": gabarito,
        "students_count": len(turma_students),
        "present_count": len(present_turma),
        "avg_total": round(sum(s["total_score"] for s in present_turma) / len(present_turma), 2) if present_turma else 0,
        "avg_reading": round(sum(s["reading_score"] for s in present_turma) / len(present_turma), 2) if present_turma else 0,
        "avg_listening": round(sum(s["listening_score"] for s in present_turma) / len(present_turma), 2) if present_turma else 0,
        "top3": [s["name"] for s in present_turma[:3]]
    }

# Overall school rank
present_all = [s for s in all_students if s["present"]]
present_all.sort(key=lambda s: (s["total_score"], s["reading_score"], s["listening_score"]), reverse=True)
current_rank = 1
for i, s in enumerate(present_all):
    if i > 0:
        prev = present_all[i-1]
        if (s["total_score"], s["reading_score"], s["listening_score"]) < (prev["total_score"], prev["reading_score"], prev["listening_score"]):
            current_rank = i + 1
    s["rank_escola"] = current_rank

for s in all_students:
    if not s["present"]:
        s["rank_escola"] = "-"
        
with open("portal_data.json", "w", encoding="utf-8") as f:
    json.dump({
        "turmas": turmas_data,
        "students": all_students
    }, f, ensure_ascii=False, indent=2)

print(f"Extracted {len(all_students)} students total across {len(turmas_data)} turmas.")
