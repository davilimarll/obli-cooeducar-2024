import openpyxl, json

wb = openpyxl.load_workbook('Gabarito_Maratona_30_Questoes.xlsx', data_only=True)
ws = wb['Todas as Turmas']

turma_blocks = [
    ("6º Ano", 4, 18),
    ("7º Ano", 20, 42),
    ("8º Ano", 44, 53),
    ("9º Ano", 55, 66),
    ("1ª Série EM", 68, 78),
    ("2ª Série EM", 80, 92),
    ("3ª Série EM", 94, 108),
]

all_data = {}

for turma_name, start_r, end_r in turma_blocks:
    # Gabarito is at start_r + 2
    gab_row = start_r + 2
    gabarito = [ws.cell(gab_row, c).value for c in range(3, 33)]
    
    students = []
    for r in range(start_r + 3, end_r + 1):
        num = ws.cell(r, 1).value
        name = ws.cell(r, 2).value
        if not name:
            continue
        answers = [ws.cell(r, c).value for c in range(3, 33)]
        
        # Check presence: if all answers are None, student was absent
        present = any(a is not None and str(a).strip() != '' for a in answers)
        
        # Calculate scores
        # Q1-Q20: Reading, Q21-Q30: Listening
        # In 8º and 9º ano, Q4 is annulled (*), so if present, +1 pt
        r_score = 0
        l_score = 0
        details = []
        for i in range(30):
            q_num = i + 1
            correct_ans = gabarito[i]
            stu_ans = answers[i] if answers[i] is not None else ""
            stu_ans_str = str(stu_ans).strip().upper()
            
            is_annulled = (turma_name in ["8º Ano", "9º Ano"] and q_num == 4)
            is_correct = False
            
            if present:
                if is_annulled:
                    is_correct = True
                elif stu_ans_str and correct_ans and stu_ans_str == str(correct_ans).strip().upper():
                    is_correct = True
            
            if is_correct:
                if q_num <= 20:
                    r_score += 1
                else:
                    l_score += 1
                    
            details.append({
                "q": q_num,
                "domain": "Reading" if q_num <= 20 else "Listening",
                "student_ans": stu_ans_str,
                "correct_ans": "*" if is_annulled else str(correct_ans).strip().upper(),
                "is_correct": is_correct,
                "is_annulled": is_annulled
            })
            
        total_score = r_score + l_score
        pct_total = round((total_score / 30.0) * 100, 1) if present else 0.0
        pct_r = round((r_score / 20.0) * 100, 1) if present else 0.0
        pct_l = round((l_score / 10.0) * 100, 1) if present else 0.0
        
        students.append({
            "num": num,
            "name": name,
            "turma": turma_name,
            "present": present,
            "total_score": total_score,
            "pct_total": pct_total,
            "reading_score": r_score,
            "pct_reading": pct_r,
            "listening_score": l_score,
            "pct_listening": pct_l,
            "details": details
        })
        
    all_data[turma_name] = {
        "gabarito": [str(g).strip().upper() if g else "" for g in gabarito],
        "students": students
    }

print("Turmas parsed:")
for t, d in all_data.items():
    pres = sum(1 for s in d["students"] if s["present"])
    print(f" - {t}: {len(d['students'])} matriculados, {pres} presentes")
