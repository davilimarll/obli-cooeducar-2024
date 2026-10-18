import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TURMAS = [
    ("6º ANO", [
        (1, "Aisha Santos David"),
        (2, "Ana Júlia Bomfim Coutinho"),
        (3, "Arthur Vieira Trindade"),
        (4, "Benício Paim De Souza Jardim"),
        (6, "Daniel Dias Souto"),
        (7, "Davi Laender Norte Brito"),
        (8, "Davi Rocha Aguiar"),
        (9, "Dely Cotrim Batista Coelho"),
        (10, "Laura Andrade Coelho Dos Santos"),
        (11, "Laura Gomes Prates"),
        (12, "Maria Clara Almeida Coutinho"),
        (13, "Valentina Silva Soares"),
    ]),
    ("7º ANO", [
        (1, "Alice De Souza Bossi Lima"),
        (2, "Cauã Jimenez Figueiredo Costa"),
        (3, "Isabela Lima Barbosa"),
        (4, "Isadora Amaral Santos"),
        (5, "Isis Costa Rodrigues De Oliveira"),
        (6, "Isly De Souza Jardim Ferreira"),
        (7, "Iuri Emanuel Lisboa Coutinho"),
        (8, "Júlia Aguilar Vieira De Sousa Lima"),
        (9, "Kauay Francisco Almeida Honorato"),
        (10, "Lavínia Ribeiro Sena"),
        (11, "Lavínia Silva Teles"),
        (12, "Lucas Alves Silva Batista"),
        (13, "Luísa Brauer Torres Nacur Quaresma"),
        (14, "Luna Januzzi Quadros"),
        (15, "Maria Cecília Rebouças"),
        (16, "Maria Clara Nogueira Tinoco"),
        (17, "Raíssa Soares Rodrigues"),
        (18, "Sâmya Ramos Silva"),
        (19, "Sophia Luanny de Jesus Silva"),
        (20, "Tito Teles Machado"),
    ]),
    ("8º ANO", [
        (1, "Davi Sedlmayer Jardim Lauar"),
        (2, "Isadora Quaresma Silveira"),
        (3, "Jolie Meireles Silva"),
        (4, "Miguel Bento Gomes Silva"),
        (5, "Ruan Batista Neves"),
        (6, "Tiago De Figueiredo Martins Filho"),
        (7, "Vicente Soares Almeida Jardim"),
    ]),
    ("9º ANO", [
        (1, "Henry Alves Franco"),
        (2, "Isa Nascimento Almeida"),
        (3, "Isadora Lima Barbosa"),
        (4, "Isaque Ferreira e Santos"),
        (5, "Isis Fonseca Batista"),
        (7, "João Antônio Santana De Almeida"),
        (8, "Júlia Vitória Amaral Soares"),
        (9, "Pedro Gabriel Dutra Souza"),
        (10, "Sophie De Deus Amantes"),
    ]),
    ("1ª SÉRIE", [
        (1, "Ana Beatriz Soares Norte"),
        (2, "Aryane Alves Silva"),
        (3, "Bruno De Souza Tomich"),
        (4, "Lisi Mendes Rodrigues Trindade"),
        (5, "Luis Felipe Lopes Fonte Nogueira"),
        (6, "Maria Clara Santos Boa Sorte"),
        (7, "Maysa Reis De Oliveira"),
        (8, "Yasmin Rodrigues de Souza Vieira"),
    ]),
    ("2ª SÉRIE", [
        (1, "Anne Beatrice Oliveira Costa"),
        (2, "Eros Cantão Cangussu"),
        (3, "Isa Silva Nogueira De Oliveira"),
        (4, "João Gabriel Dias Souto"),
        (5, "João Pedro Aguilar Vieira Braga"),
        (6, "Luiz Fernando Alves Da Silva Filho"),
        (7, "Mariana Quaresma Silveira"),
        (8, "Melissa Silva Cardoso"),
        (9, "Sam Antunes Cabral"),
        (10, "Sara Barbosa Da Silva"),
    ]),
    ("3ª SÉRIE", [
        (1, "Allicia Sena De Figueiredo"),
        (2, "Ana Maria Ferreira Morais"),
        (3, "Breno Martins Coelho Da Cruz"),
        (4, "Davy Luca Galvao Da Silva"),
        (5, "João Gabriel Lemes Almeida Dos Santos"),
        (6, "Joaquim Amaral Mendes"),
        (7, "Kevyn Leonardo Ramos Rodrigues"),
        (8, "Laura Lima Porto"),
        (9, "Maria Eduarda Rebouças Brandão"),
        (10, "Maria Laura Alves Luz"),
        (11, "Samuel Sedlmayer Souza"),
        (12, "Vitor José Rodrigues Jardim"),
    ]),
]

# Paleta de cores elegante
NAVY_DARK = "1B365D"
NAVY_LIGHT = "2E5B88"
TURMA_BANNER_BG = "0F172A"
GOLD_BG = "FEF3C7"
GOLD_TEXT = "92400E"
ZEBRA_LGT = "F8FAFC"
BORDER_CLR = "CBD5E1"
SCORE_HDR_BG = "1E3A8A"

font_title = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="E2E8F0")
font_turma = Font(name="Segoe UI", size=12, bold=True, color="FDE68A")
font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
font_official = Font(name="Segoe UI", size=11, bold=True, color=GOLD_TEXT)
font_student_name = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
font_student_num = Font(name="Segoe UI", size=10, bold=True, color="475569")
font_cell = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
font_score = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")

fill_title = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
fill_turma = PatternFill(start_color=TURMA_BANNER_BG, end_color=TURMA_BANNER_BG, fill_type="solid")
fill_header = PatternFill(start_color=NAVY_LIGHT, end_color=NAVY_LIGHT, fill_type="solid")
fill_official = PatternFill(start_color=GOLD_BG, end_color=GOLD_BG, fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
fill_zebra = PatternFill(start_color=ZEBRA_LGT, end_color=ZEBRA_LGT, fill_type="solid")
fill_score_hdr = PatternFill(start_color=SCORE_HDR_BG, end_color=SCORE_HDR_BG, fill_type="solid")

thin_side = Side(border_style="thin", color=BORDER_CLR)
border_box = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
border_official = Border(
    left=thin_side, right=thin_side, top=thin_side,
    bottom=Side(border_style="medium", color="D97706")
)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center", indent=1)


def set_column_widths(ws):
    ws.column_dimensions["A"].width = 6      # Nº
    ws.column_dimensions["B"].width = 36     # Nome do Aluno
    for q in range(1, 31):
        col_letter = get_column_letter(2 + q)
        ws.column_dimensions[col_letter].width = 4.6  # Quadradinho perfeito
    ws.column_dimensions["AG"].width = 14    # Acertos (30)
    ws.column_dimensions["AH"].width = 12    # % Acertos


def write_turma_block(ws, start_row, turma_nome, alunos, show_banner=True):
    r = start_row
    if show_banner:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=34)
        b_cell = ws.cell(row=r, column=1, value=f"TURMA: {turma_nome} ({len(alunos)} alunos)")
        b_cell.font = font_turma
        b_cell.fill = fill_turma
        b_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 28
        r += 1

    # Linha de Cabeçalhos das Colunas
    headers = ["Nº", "Nome do Aluno"] + [f"Q{i:02d}" for i in range(1, 31)] + ["Acertos (30)", "% Acertos"]
    ws.row_dimensions[r].height = 26
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=col_idx, value=h)
        cell.font = font_header
        cell.alignment = align_center
        cell.fill = fill_score_hdr if col_idx >= 33 else fill_header
        cell.border = border_box
    r += 1

    # Linha do GABARITO OFICIAL daquela turma
    ofic_row = r
    ws.row_dimensions[ofic_row].height = 28
    c_num = ws.cell(row=ofic_row, column=1, value="★")
    c_num.font = font_official
    c_num.fill = fill_official
    c_num.alignment = align_center
    c_num.border = border_official

    c_name = ws.cell(row=ofic_row, column=2, value=f"[ GABARITO OFICIAL — {turma_nome} ]")
    c_name.font = font_official
    c_name.fill = fill_official
    c_name.alignment = align_left
    c_name.border = border_official

    for q in range(1, 31):
        col_idx = 2 + q
        c_q = ws.cell(row=ofic_row, column=col_idx, value="")
        c_q.font = font_official
        c_q.fill = fill_official
        c_q.alignment = align_center
        c_q.border = border_official

    c_tot_ofic = ws.cell(row=ofic_row, column=33, value=f"=COUNTA(C{ofic_row}:AF{ofic_row})")
    c_tot_ofic.font = Font(name="Segoe UI", size=9, bold=True, color=GOLD_TEXT)
    c_tot_ofic.fill = fill_official
    c_tot_ofic.alignment = align_center
    c_tot_ofic.border = border_official

    c_pct_ofic = ws.cell(row=ofic_row, column=34, value="Gabarito")
    c_pct_ofic.font = Font(name="Segoe UI", size=9, bold=True, color=GOLD_TEXT)
    c_pct_ofic.fill = fill_official
    c_pct_ofic.alignment = align_center
    c_pct_ofic.border = border_official
    r += 1

    # Linhas dos Alunos
    for idx, (num, nome) in enumerate(alunos):
        ws.row_dimensions[r].height = 25
        row_fill = fill_zebra if (idx % 2 == 1) else fill_white

        c_n = ws.cell(row=r, column=1, value=num)
        c_n.font = font_student_num
        c_n.alignment = align_center
        c_n.fill = row_fill
        c_n.border = border_box

        c_nome = ws.cell(row=r, column=2, value=nome)
        c_nome.font = font_student_name
        c_nome.alignment = align_left
        c_nome.fill = row_fill
        c_nome.border = border_box

        for q in range(1, 31):
            col_idx = 2 + q
            cq = ws.cell(row=r, column=col_idx, value="")
            cq.font = font_cell
            cq.alignment = align_center
            cq.fill = row_fill
            cq.border = border_box

        # Fórmula automática de Acertos comparando com o gabarito da turma (ofic_row)
        c_acertos = ws.cell(row=r, column=33)
        c_acertos.value = f'=SUMPRODUCT((C{r}:AF{r}<>"")*(C{r}:AF{r}=$C${ofic_row}:$AF${ofic_row})*($C${ofic_row}:$AF${ofic_row}<>""))'
        c_acertos.font = font_score
        c_acertos.alignment = align_center
        c_acertos.fill = row_fill
        c_acertos.border = border_box

        c_pct = ws.cell(row=r, column=34)
        c_pct.value = f'=IF(COUNTA($C${ofic_row}:$AF${ofic_row})>0, AG{r}/COUNTA($C${ofic_row}:$AF${ofic_row}), 0)'
        c_pct.number_format = '0.0%'
        c_pct.font = font_score
        c_pct.alignment = align_center
        c_pct.fill = row_fill
        c_pct.border = border_box

        r += 1

    return r


def create_maratona_sheet():
    wb = openpyxl.Workbook()

    # =========================================================
    # ABA 1: TODAS AS TURMAS (UNIFICADA)
    # =========================================================
    ws_all = wb.active
    ws_all.title = "Todas as Turmas"
    ws_all.views.sheetView[0].showGridLines = True

    ws_all.merge_cells("A1:AH1")
    t_cell = ws_all["A1"]
    t_cell.value = "MARATONA 2026 — GABARITO DE TODAS AS TURMAS (6º ANO À 3ª SÉRIE — 30 QUESTÕES)"
    t_cell.font = font_title
    t_cell.fill = fill_title
    t_cell.alignment = align_center
    ws_all.row_dimensions[1].height = 36

    ws_all.merge_cells("A2:AH2")
    s_cell = ws_all["A2"]
    s_cell.value = "Preencha a alternativa (A, B, C, D, E) nos quadradinhos de cada aluno. O Gabarito Oficial de cada turma calcula os acertos na hora!"
    s_cell.font = font_subtitle
    s_cell.fill = fill_title
    s_cell.alignment = align_center
    ws_all.row_dimensions[2].height = 22

    curr_row = 4
    for turma_nome, alunos in TURMAS:
        curr_row = write_turma_block(ws_all, curr_row, turma_nome, alunos, show_banner=True)
        curr_row += 1  # Linha em branco separadora entre turmas

    set_column_widths(ws_all)

    # =========================================================
    # ABAS INDIVIDUAIS POR TURMA (PARA NAVEGAÇÃO RÁPIDA SE PREFERIR)
    # =========================================================
    for turma_nome, alunos in TURMAS:
        ws_t = wb.create_sheet(title=turma_nome)
        ws_t.views.sheetView[0].showGridLines = True

        ws_t.merge_cells("A1:AH1")
        tc = ws_t["A1"]
        tc.value = f"MARATONA 2026 — FOLHA DE RESPOSTAS ({turma_nome})"
        tc.font = font_title
        tc.fill = fill_title
        tc.alignment = align_center
        ws_t.row_dimensions[1].height = 34

        ws_t.merge_cells("A2:AH2")
        sc = ws_t["A2"]
        sc.value = "Preencha as letras (A, B, C, D, E) nos quadradinhos abaixo. Os acertos são calculados automaticamente pelo Gabarito Oficial na linha 4."
        sc.font = font_subtitle
        sc.fill = fill_title
        sc.alignment = align_center
        ws_t.row_dimensions[2].height = 20

        write_turma_block(ws_t, 3, turma_nome, alunos, show_banner=False)
        set_column_widths(ws_t)

    output_path = "Gabarito_Maratona_30_Questoes.xlsx"
    wb.save(output_path)
    print(f"Planilha completa gerada com sucesso em: {output_path}")


if __name__ == "__main__":
    create_maratona_sheet()
